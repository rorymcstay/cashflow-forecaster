"""Generates household_budget_template.xlsx"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

CURRENCY = '£#,##0.00'

# ---- Styles ----
HEADER_FILL = PatternFill('solid', fgColor='2F5233')
SECTION_FILL = PatternFill('solid', fgColor='E8EDE9')
TOTAL_FILL = PatternFill('solid', fgColor='D9E4DD')
INPUT_FILL = PatternFill('solid', fgColor='FFF9E6')
HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
SECTION_FONT = Font(bold=True, size=10, color='2F5233')
TOTAL_FONT = Font(bold=True, size=10)
TITLE_FONT = Font(bold=True, size=16, color='2F5233')
SUBTITLE_FONT = Font(italic=True, size=10, color='666666')
thin = Side(style='thin', color='CCCCCC')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM_COLS_BUDGET = 5  # Category, Item, Monthly Amount, Day of Month, Notes

wb = openpyxl.Workbook()

# =====================================================================
# SHEET 1: BUDGET
# =====================================================================
ws = wb.active
ws.title = 'Budget'
ws.sheet_view.showGridLines = False

ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 24
ws.column_dimensions['C'].width = 16
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 30

ws['A1'] = 'Household Budget'
ws['A1'].font = TITLE_FONT
ws['A2'] = 'Fill in the yellow cells. Day of Month is optional — set it to feed the Cash Flow Forecast sheet.'
ws['A2'].font = SUBTITLE_FONT

row = 4

def header_row(r, cols=('Category', 'Item', 'Monthly Amount', 'Day of Month', 'Notes')):
    for i, c in enumerate(cols):
        cell = ws.cell(row=r, column=1 + i, value=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = BORDER
    ws.row_dimensions[r].height = 20

def section_row(r, label):
    cell = ws.cell(row=r, column=1, value=label)
    cell.font = SECTION_FONT
    for col in range(1, NUM_COLS_BUDGET + 1):
        ws.cell(row=r, column=col).fill = SECTION_FILL
        ws.cell(row=r, column=col).border = BORDER

def item_row(r, category, item, amount, day=None, note='', editable_amount=True):
    ws.cell(row=r, column=1, value=category).border = BORDER
    ws.cell(row=r, column=2, value=item).border = BORDER
    acell = ws.cell(row=r, column=3, value=amount)
    acell.number_format = CURRENCY
    acell.border = BORDER
    if editable_amount:
        acell.fill = INPUT_FILL
    dcell = ws.cell(row=r, column=4, value=day)
    dcell.border = BORDER
    dcell.fill = INPUT_FILL
    dcell.alignment = Alignment(horizontal='center')
    day_dv.add(dcell.coordinate)
    ncell = ws.cell(row=r, column=5, value=note)
    ncell.border = BORDER
    ncell.font = Font(italic=True, size=9, color='888888')

def subtotal_row(r, label, first_data_row, last_data_row):
    ws.cell(row=r, column=1, value=label).font = TOTAL_FONT
    for col in range(1, NUM_COLS_BUDGET + 1):
        ws.cell(row=r, column=col).fill = TOTAL_FILL
        ws.cell(row=r, column=col).border = BORDER
    cell = ws.cell(row=r, column=3, value=f'=SUM(C{first_data_row}:C{last_data_row})')
    cell.number_format = CURRENCY
    cell.font = TOTAL_FONT
    return r

# Day-of-month validation (1-31, optional)
day_dv = DataValidation(type='whole', operator='between', formula1=1, formula2=31,
                         allow_blank=True, showErrorMessage=True,
                         errorTitle='Invalid day', error='Enter a day of month between 1 and 31, or leave blank.')
ws.add_data_validation(day_dv)

header_row(row)
row += 1

# --- INCOME ---
section_row(row, 'INCOME'); row += 1
income_start = row
item_row(row, 'Income', 'Household member 1 (take-home)', 0, note='Fill in'); row += 1
item_row(row, 'Income', 'Household member 2 (take-home)', 0, note='Fill in'); row += 1
item_row(row, 'Income', 'Other income', 0, note='e.g. side income, bonuses'); row += 1
# a few blank editable rows for additional income sources
for _ in range(3):
    item_row(row, 'Income', '', 0); row += 1
income_end = row - 1
income_total_row = row
subtotal_row(row, 'Total Income', income_start, income_end); row += 2

# --- HOUSING ---
section_row(row, 'HOUSING'); row += 1
housing_start = row
item_row(row, 'Housing', 'Rent / Mortgage', 2800, day=17,
         note='Joint account — TLP RE Client Acc (letting agent), confirmed 12/12 months'); row += 1
item_row(row, 'Housing', 'Other housing costs', 0,
         note='Council Tax is paid as an annual lump sum — see Upcoming Expenses (~£1,020, due ~Apr)'); row += 1
housing_end = row - 1
housing_total_row = row
subtotal_row(row, 'Total Housing', housing_start, housing_end); row += 2

# --- UTILITIES & INSURANCE ---
section_row(row, 'UTILITIES & INSURANCE'); row += 1
util_start = row
item_row(row, 'Utilities', 'Electricity', 68.76, day=1,
         note='E.ON Next DD, Joint account — actual avg £68.76 (range £62-76), was budgeted £64'); row += 1
item_row(row, 'Utilities', 'WiFi', 20.56, day=5,
         note='Community Fibre DD, Joint account — confirmed, very stable'); row += 1
item_row(row, 'Insurance', 'Home insurance', 23,
         note='Not found in reviewed statements — verify still active / correct amount'); row += 1
item_row(row, 'Insurance', 'Vet insurance', 23.02, day=9,
         note='Urban Jungle DD, Joint account — renewed from £15.02 to £23.02/mo in Apr 2026'); row += 1
util_end = row - 1
util_total_row = row
subtotal_row(row, 'Total Utilities & Insurance', util_start, util_end); row += 2

# --- TRANSPORT ---
section_row(row, 'TRANSPORT'); row += 1
transport_start = row
item_row(row, 'Transport', 'Car tax — B355RHC (DVLA)', 31.50, day=1,
         note='Personal HSBC DD — was bundled into "Car expenses 142"'); row += 1
item_row(row, 'Transport', 'Car tax — C800BWR (DVLA)', 32.81, day=1,
         note='Personal HSBC DD — second car, was bundled into "Car expenses 142"'); row += 1
item_row(row, 'Transport', 'Car insurance (Elephant)', 77.53, day=28,
         note='Personal HSBC DD — was bundled into "Car expenses 142"'); row += 1
item_row(row, 'Transport', 'Fuel', 102.54,
         note='Variable card spend, Joint account — avg per active month, not a fixed DD (was "Car expenses 100")'); row += 1
transport_end = row - 1
transport_total_row = row
subtotal_row(row, 'Total Transport', transport_start, transport_end); row += 2

# --- SUBSCRIPTIONS ---
section_row(row, 'SUBSCRIPTIONS'); row += 1
sub_start = row
item_row(row, 'Subscriptions', 'Spotify', 17.66, day=1,
         note='Personal Monzo — confirmed, day varies 1-11'); row += 1
item_row(row, 'Subscriptions', 'Netflix', 5.99, day=5,
         note='Joint account — confirmed exact'); row += 1
item_row(row, 'Subscriptions', 'Amazon Prime', 8.99, day=16,
         note='Amex — confirmed exact'); row += 1
item_row(row, 'Subscriptions', 'Apple', 3,
         note='Not confirmed monthly — only a one-off £53.99 Amex charge seen; verify what this is'); row += 1
item_row(row, 'Subscriptions', 'Google', 2,
         note='Not found in reviewed statements — verify still active'); row += 1
item_row(row, 'Subscriptions', 'TV licence', 36.00, day=1,
         note='TV Licensing DD, Joint account — actual is £36, was budgeted £15; only seen Jun & Jul (may be a 10-month DD scheme)'); row += 1
item_row(row, 'Subscriptions', 'AWS', 23.53,
         note='Usage-based, Personal Monzo — avg £23.53 (range ~£14-33), varies month to month, was budgeted £5'); row += 1
item_row(row, 'Subscriptions', 'Claude.ai', 18.00, day=28,
         note='New — Amex, confirmed 3/3 months, not in original list'); row += 1
item_row(row, 'Subscriptions', 'Voxi (Vodafone) mobile', 10.00, day=8,
         note='New — Personal HSBC, variable £4.80-16, present every month, not in original list'); row += 1
sub_end = row - 1
sub_total_row = row
subtotal_row(row, 'Total Subscriptions', sub_start, sub_end); row += 2

# --- PERSONAL & OTHER ---
section_row(row, 'PERSONAL & OTHER'); row += 1
personal_start = row
item_row(row, 'Personal', 'Nathan', 150,
         note='Not found in reviewed statements — verify (may be paid via cash/another account)'); row += 1
item_row(row, 'Personal', 'Swimming', 45,
         note='Not found in reviewed statements — verify (may be paid via cash/another account)'); row += 1
item_row(row, 'Personal', 'The Riders Hub', 15.00,
         note='Low confidence — appears sporadically on personal HSBC & Monzo, possibly a lesson/hobby subscription; verify'); row += 1
personal_end = row - 1
personal_total_row = row
subtotal_row(row, 'Total Personal & Other', personal_start, personal_end); row += 2

# --- SAVINGS & INVESTMENTS CONTRIBUTION ---
section_row(row, 'SAVINGS & INVESTMENTS'); row += 1
savings_start = row
item_row(row, 'Savings', 'Monthly savings contribution', 0, note='Fill in'); row += 1
item_row(row, 'Investments', 'Monthly investment contribution', 0, note='Fill in'); row += 1
savings_end = row - 1
savings_total_row = row
subtotal_row(row, 'Total Savings & Investments', savings_start, savings_end); row += 2

# --- SUMMARY ---
section_row(row, 'SUMMARY'); row += 1
ws.cell(row=row, column=1, value='Total Income').font = TOTAL_FONT
c = ws.cell(row=row, column=3, value=f'=C{income_total_row}')
c.number_format = CURRENCY
for col in range(1, NUM_COLS_BUDGET + 1): ws.cell(row=row, column=col).border = BORDER
row += 1

expense_rows = [housing_total_row, util_total_row, transport_total_row, sub_total_row, personal_total_row, savings_total_row]
ws.cell(row=row, column=1, value='Total Outgoings (Housing + Expenses + Savings)').font = TOTAL_FONT
c = ws.cell(row=row, column=3, value='=' + '+'.join(f'C{r}' for r in expense_rows))
c.number_format = CURRENCY
for col in range(1, NUM_COLS_BUDGET + 1): ws.cell(row=row, column=col).border = BORDER
outgoings_row = row
row += 1

ws.cell(row=row, column=1, value='Net Remaining').font = Font(bold=True, size=12)
c = ws.cell(row=row, column=3, value=f'=C{income_total_row}-C{outgoings_row}')
c.number_format = CURRENCY
c.font = Font(bold=True, size=12)
for col in range(1, NUM_COLS_BUDGET + 1):
    ws.cell(row=row, column=col).fill = TOTAL_FILL
    ws.cell(row=row, column=col).border = BORDER
row += 2

ws.freeze_panes = 'A5'
BUDGET_LAST_ROW = row  # generous upper bound used by Cash Flow SUMIFS ranges

# =====================================================================
# SHEET 2: SAVINGS & INVESTMENTS
# =====================================================================
ws2 = wb.create_sheet('Savings & Investments')
ws2.sheet_view.showGridLines = False
for col, w in zip('ABCDEFG', [22, 16, 16, 18, 16, 16, 28]):
    ws2.column_dimensions[col].width = w

ws2['A1'] = 'Savings & Investments Tracker'
ws2['A1'].font = TITLE_FONT
ws2['A2'] = 'One row per account/pot. Fill in the yellow cells; totals and progress update automatically.'
ws2['A2'].font = SUBTITLE_FONT

headers = ['Account / Pot', 'Type', 'Starting Balance', 'Monthly Contribution', 'Current Balance', 'Goal Target', 'Notes']
hr = 4
for i, h in enumerate(headers):
    cell = ws2.cell(row=hr, column=1 + i, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal='left', vertical='center')
ws2.row_dimensions[hr].height = 20

data_start = hr + 1
example_rows = [
    ('Emergency Fund', 'Savings', 0, 0, 0, 0, ''),
    ('Stocks & Shares ISA', 'Investment', 0, 0, 0, 0, ''),
    ('Pension', 'Investment', 0, 0, 0, 0, ''),
]
r = data_start
for name, typ, start_bal, contrib, cur_bal, goal, note in example_rows:
    ws2.cell(row=r, column=1, value=name).border = BORDER
    tcell = ws2.cell(row=r, column=2, value=typ)
    tcell.border = BORDER
    for col, val in zip((3, 4, 5, 6), (start_bal, contrib, cur_bal, goal)):
        c = ws2.cell(row=r, column=col, value=val)
        c.number_format = CURRENCY
        c.border = BORDER
        c.fill = INPUT_FILL
    nc = ws2.cell(row=r, column=7, value=note)
    nc.border = BORDER
    r += 1
data_end = r - 1

# a few blank editable rows for the user to extend
for _ in range(4):
    ws2.cell(row=r, column=1).border = BORDER
    ws2.cell(row=r, column=2).border = BORDER
    for col in (3, 4, 5, 6):
        c = ws2.cell(row=r, column=col)
        c.number_format = CURRENCY
        c.border = BORDER
        c.fill = INPUT_FILL
    ws2.cell(row=r, column=7).border = BORDER
    r += 1
data_end_extended = r - 1

# Totals row
total_row = r + 1
ws2.cell(row=total_row, column=1, value='TOTAL').font = TOTAL_FONT
for col in (1, 2, 7):
    ws2.cell(row=total_row, column=col).fill = TOTAL_FILL
    ws2.cell(row=total_row, column=col).border = BORDER
for col in (3, 4, 5, 6):
    colletter = get_column_letter(col)
    c = ws2.cell(row=total_row, column=col, value=f'=SUM({colletter}{data_start}:{colletter}{data_end_extended})')
    c.number_format = CURRENCY
    c.font = TOTAL_FONT
    c.fill = TOTAL_FILL
    c.border = BORDER

dv2 = DataValidation(type='list', formula1='"Savings,Investment,Pension,Other"', allow_blank=True)
ws2.add_data_validation(dv2)
dv2.add(f'B{data_start}:B{data_end_extended}')

ws2.freeze_panes = 'A5'

# =====================================================================
# SHEET 3: UPCOMING EXPENSES
# =====================================================================
ws_ue = wb.create_sheet('Upcoming Expenses')
ws_ue.sheet_view.showGridLines = False
for col, w in zip('ABCD', [16, 34, 14, 30]):
    ws_ue.column_dimensions[col].width = w

ws_ue['A1'] = 'Upcoming Expenses'
ws_ue['A1'].font = TITLE_FONT
ws_ue['A2'] = ('One-off items that crop up as and when they come up (e.g. car MOT, a gift, a repair). '
               'Enter Amount as a positive cost — it is automatically treated as money going out. '
               'Any item whose date falls in the current Cash Flow Forecast month is pulled in automatically.')
ws_ue['A2'].font = SUBTITLE_FONT

ue_headers = ['Date', 'Description', 'Amount', 'Notes']
UE_HEADER_ROW = 4
for i, h in enumerate(ue_headers):
    cell = ws_ue.cell(row=UE_HEADER_ROW, column=1 + i, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='left', vertical='center')
    cell.border = BORDER
ws_ue.row_dimensions[UE_HEADER_ROW].height = 20

UE_FIRST_ROW = UE_HEADER_ROW + 1
UE_ROWS = 60
UE_LAST_ROW = UE_FIRST_ROW + UE_ROWS - 1

# Known non-monthly bills found in the statement history — paid too infrequently
# to belong in the monthly Budget sheet, but real and worth forecasting for.
import datetime as _dt
known_upcoming = [
    (_dt.date(2026, 10, 14), 'Thames Water (6-monthly)', 188.00,
     'Recurs ~every 6 months; last seen ~Apr 2026 (~£186-190), Joint account'),
    (_dt.date(2027, 4, 27), 'Council Tax (Wandsworth) — annual', 1020.35,
     'Paid as an annual lump sum from the Joint account; may increase with rate changes'),
]

for i in range(UE_ROWS):
    r = UE_FIRST_ROW + i
    dcell = ws_ue.cell(row=r, column=1)
    dcell.number_format = 'DD-MMM-YYYY'
    dcell.fill = INPUT_FILL
    dcell.border = BORDER
    ws_ue.cell(row=r, column=2).border = BORDER
    ws_ue.cell(row=r, column=2).fill = INPUT_FILL
    acell = ws_ue.cell(row=r, column=3)
    acell.number_format = CURRENCY
    acell.fill = INPUT_FILL
    acell.border = BORDER
    ncell = ws_ue.cell(row=r, column=4)
    ncell.fill = INPUT_FILL
    ncell.border = BORDER
    if i < len(known_upcoming):
        d, desc, amt, note = known_upcoming[i]
        dcell.value = d
        ws_ue.cell(row=r, column=2).value = desc
        acell.value = amt
        ncell.value = note

ws_ue.freeze_panes = 'A5'

# =====================================================================
# SHEET 4: CASH FLOW FORECAST
# =====================================================================
ws3 = wb.create_sheet('Cash Flow Forecast')
ws3.sheet_view.showGridLines = False
for col, w in zip('ABCDEFGHI', [13, 13, 13, 13, 13, 14, 13, 13, 14]):
    ws3.column_dimensions[col].width = w

ws3['A1'] = 'Cash Flow Forecast'
ws3['A1'].font = TITLE_FONT
ws3['A2'] = ('Set Forecast Month + Opening Balance below. Realised Net overrides Forecast Net for a day when filled in. '
             'Amounts are pulled from the Budget sheet by Day of Month, plus any items from Upcoming Expenses that fall in this month.')
ws3['A2'].font = SUBTITLE_FONT

# --- Inputs ---
ws3['A4'] = 'Forecast Month:'
ws3['A4'].font = TOTAL_FONT
mcell = ws3['B4']
mcell.value = __import__('datetime').date.today().replace(day=1)
mcell.number_format = 'MMMM YYYY'
mcell.fill = INPUT_FILL
mcell.border = BORDER
mcell.font = Font(bold=True)

ws3['D4'] = 'Opening Balance:'
ws3['D4'].font = TOTAL_FONT
ocell = ws3['E4']
ocell.value = 0
ocell.number_format = CURRENCY
ocell.fill = INPUT_FILL
ocell.border = BORDER
ocell.font = Font(bold=True)

wb.defined_names['ForecastMonth'] = DefinedName('ForecastMonth', attr_text="'Cash Flow Forecast'!$B$4")
wb.defined_names['OpeningBalance'] = DefinedName('OpeningBalance', attr_text="'Cash Flow Forecast'!$E$4")

# --- Daily table header (row 6) ---
DAY_HEADER_ROW = 6
daily_headers = ['Date', 'Forecast In', 'Forecast Out', 'Upcoming\nExpenses', 'Forecast Net',
                  'Forecast\nBalance', 'Realised Net\n(manual)', 'Effective Net', 'Effective\nBalance']
for i, h in enumerate(daily_headers):
    cell = ws3.cell(row=DAY_HEADER_ROW, column=1 + i, value=h)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = BORDER
ws3.row_dimensions[DAY_HEADER_ROW].height = 30

# --- Daily rows (up to 31 days) ---
DAY_FIRST_ROW = DAY_HEADER_ROW + 1
for n in range(1, 32):
    r = DAY_FIRST_ROW + n - 1
    a = f'A{r}'
    # Date (blank if this day doesn't exist in the forecast month)
    ws3.cell(row=r, column=1,
             value=f'=IF({n}<=DAY(EOMONTH($B$4,0)),DATE(YEAR($B$4),MONTH($B$4),{n}),"")')
    ws3.cell(row=r, column=1).number_format = 'DD-MMM (DDD)'

    # Forecast Inflow / Outflow from Budget sheet, matched by Day of Month
    ws3.cell(row=r, column=2,
             value=f'=IF({a}="","",SUMIFS(Budget!$C$5:$C${BUDGET_LAST_ROW},'
                   f'Budget!$A$5:$A${BUDGET_LAST_ROW},"Income",Budget!$D$5:$D${BUDGET_LAST_ROW},{n}))')
    ws3.cell(row=r, column=3,
             value=f'=IF({a}="","",SUMIFS(Budget!$C$5:$C${BUDGET_LAST_ROW},'
                   f'Budget!$A$5:$A${BUDGET_LAST_ROW},"<>Income",Budget!$D$5:$D${BUDGET_LAST_ROW},{n}))')

    # Upcoming Expenses falling on this date (entered as positive costs, applied as an outflow)
    ws3.cell(row=r, column=4,
             value=f'=IF({a}="","",-SUMIFS(\'Upcoming Expenses\'!$C${UE_FIRST_ROW}:$C${UE_LAST_ROW},'
                   f'\'Upcoming Expenses\'!$A${UE_FIRST_ROW}:$A${UE_LAST_ROW},{a}))')

    # Forecast Net
    ws3.cell(row=r, column=5, value=f'=IF({a}="","",B{r}-C{r}+D{r})')

    # Forecast running Balance
    if n == 1:
        ws3.cell(row=r, column=6, value=f'=IF({a}="","",$E$4+E{r})')
    else:
        ws3.cell(row=r, column=6, value=f'=IF({a}="","",F{r - 1}+E{r})')

    # Realised Net (manual entry, blank = not yet realised)
    gcell = ws3.cell(row=r, column=7)
    gcell.fill = INPUT_FILL

    # Effective Net = Realised if present, else Forecast
    ws3.cell(row=r, column=8, value=f'=IF({a}="","",IF(G{r}<>"",G{r},E{r}))')

    # Effective running Balance
    if n == 1:
        ws3.cell(row=r, column=9, value=f'=IF({a}="","",$E$4+H{r})')
    else:
        ws3.cell(row=r, column=9, value=f'=IF({a}="","",I{r - 1}+H{r})')

    for col in range(1, 10):
        cell = ws3.cell(row=r, column=col)
        cell.border = BORDER
        if col in (2, 3, 4, 5, 6, 7, 8, 9):
            cell.number_format = CURRENCY
        if col == 1:
            cell.alignment = Alignment(horizontal='center')

DAY_LAST_ROW = DAY_FIRST_ROW + 31 - 1

# --- Closing balance summary ---
summary_row = DAY_LAST_ROW + 2
ws3.cell(row=summary_row, column=1, value='Month Closing Balance (Forecast):').font = TOTAL_FONT
cbf = ws3.cell(row=summary_row, column=3,
               value=f'=LOOKUP(2,1/($A${DAY_FIRST_ROW}:$A${DAY_LAST_ROW}<>""),$F${DAY_FIRST_ROW}:$F${DAY_LAST_ROW})')
cbf.number_format = CURRENCY
cbf.font = TOTAL_FONT

ws3.cell(row=summary_row + 1, column=1, value='Month Closing Balance (Effective):').font = TOTAL_FONT
cbe = ws3.cell(row=summary_row + 1, column=3,
               value=f'=LOOKUP(2,1/($A${DAY_FIRST_ROW}:$A${DAY_LAST_ROW}<>""),$I${DAY_FIRST_ROW}:$I${DAY_LAST_ROW})')
cbe.number_format = CURRENCY
cbe.font = TOTAL_FONT
for rr in (summary_row, summary_row + 1):
    for col in (1, 3):
        ws3.cell(row=rr, column=col).fill = TOTAL_FILL
        ws3.cell(row=rr, column=col).border = BORDER

wb.defined_names['ClosingEffectiveBalance'] = DefinedName(
    'ClosingEffectiveBalance', attr_text=f"'Cash Flow Forecast'!$C${summary_row + 1}")

ws3.freeze_panes = 'A7'

wb.save('household_budget_template.xlsx')
print('Saved household_budget_template.xlsx')
print(f'BUDGET_LAST_ROW={BUDGET_LAST_ROW}  UE_FIRST_ROW={UE_FIRST_ROW}  UE_LAST_ROW={UE_LAST_ROW}  '
      f'DAY_FIRST_ROW={DAY_FIRST_ROW}  DAY_LAST_ROW={DAY_LAST_ROW}')
